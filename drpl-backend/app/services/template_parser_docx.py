"""
DRPL Backend - DOCX Template Parser
Extracts structure from uploaded Word documents and uses AI to parse the template.
"""

import json
import logging
from typing import Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.services.ai_service import call_ai

logger = logging.getLogger(__name__)


def _extract_docx_structure(file_path: str) -> dict:
    """Extract heading hierarchy, paragraphs, tables, and placeholders from a DOCX file."""
    doc = Document(file_path)

    sections = []
    current_section = None
    body_text = []

    for para in doc.paragraphs:
        style_name = (para.style.name or "").lower()
        text = para.text.strip()
        if not text:
            continue

        # Detect headings
        if style_name.startswith("heading"):
            level = 1
            for ch in style_name:
                if ch.isdigit():
                    level = int(ch)
                    break

            if level <= 2:
                # Start new section
                if current_section:
                    sections.append(current_section)
                current_section = {
                    "title": text,
                    "level": level,
                    "paragraphs": [],
                    "placeholders": [],
                }
            elif current_section:
                current_section["paragraphs"].append(f"[H{level}] {text}")
        elif current_section:
            current_section["paragraphs"].append(text)
            # Detect placeholders like [COMPANY_NAME] or <<FIELD>>
            import re
            placeholders = re.findall(r'\[([A-Z_]+)\]|<<([^>]+)>>|\{([A-Z_]+)\}', text)
            for match in placeholders:
                ph = next(m for m in match if m)
                if ph not in current_section["placeholders"]:
                    current_section["placeholders"].append(ph)
        else:
            body_text.append(text)

    if current_section:
        sections.append(current_section)

    # Extract tables
    tables = []
    for i, table in enumerate(doc.tables):
        if not table.rows:
            continue
        headers = [cell.text.strip() for cell in table.rows[0].cells]
        row_count = len(table.rows) - 1  # exclude header
        tables.append({
            "index": i,
            "headers": headers,
            "data_rows": row_count,
        })

    return {
        "sections": sections,
        "tables": tables,
        "preamble": body_text[:10],  # First paragraphs before any heading
        "total_sections": len(sections),
        "total_tables": len(tables),
    }


async def parse_docx_template(file_path: str, db=None) -> dict:
    """
    Extract structure from a DOCX template and use AI to produce
    a semantic analysis in the standard template JSON format.
    """
    try:
        raw = _extract_docx_structure(file_path)
    except Exception as e:
        logger.error(f"Failed to parse DOCX {file_path}: {e}")
        return {"sections": [], "metadata": {"error": str(e)}}

    if not raw["sections"] and not raw["tables"]:
        return {"sections": [], "metadata": {"error": "No structure found in DOCX"}}

    # Build a text representation for AI
    lines = []
    for s in raw["sections"]:
        lines.append(f"{'#' * s['level']} {s['title']}")
        for p in s["paragraphs"][:5]:
            lines.append(f"  {p}")
        if s["placeholders"]:
            lines.append(f"  Placeholders: {', '.join(s['placeholders'])}")
        lines.append("")

    for t in raw["tables"]:
        lines.append(f"TABLE {t['index']+1}: {' | '.join(t['headers'])} ({t['data_rows']} data rows)")

    text_repr = "\n".join(lines)
    if len(text_repr) > 15000:
        text_repr = text_repr[:15000] + "\n[... truncated ...]"

    system_prompt = """You are a document structure analyst. Given the extracted structure of a Word document template,
produce a structured JSON representation.

Return ONLY valid JSON with this structure:
{
  "sections": [
    {
      "title": "Section title",
      "description": "What this section should contain",
      "order": 1,
      "fields": [
        {"name": "field_name", "label": "Field Label", "type": "text|number|date|textarea|table", "required": true}
      ],
      "tables": [
        {"name": "table_name", "columns": ["Col1", "Col2"], "description": "What data goes here"}
      ]
    }
  ],
  "metadata": {
    "template_type": "works_single_packet|service_two_packet|supply|amc|custom",
    "document_category": "annexure|declaration|certificate|boq|letter|proposal|custom",
    "contract_type": "Works|Service|Supply|null",
    "format_rules": ["rule1", "rule2"],
    "required_sections": ["header", "body", "signature_block"],
    "match_patterns": ["pattern*"]
  },
  "content_template_markdown": "The original document converted to a markdown template with placeholders"
}

Capture every section, field, table, and placeholder. Infer document_category from content."""

    user_prompt = f"Parse this DOCX template structure:\n\n{text_repr}"

    try:
        result = await call_ai(system_prompt, user_prompt, db, "template_parser")
        result = result.strip()
        if result.startswith("```"):
            result = result.split("\n", 1)[1].rsplit("```", 1)[0]
        structure = json.loads(result)
        return structure
    except json.JSONDecodeError as e:
        logger.error(f"AI returned invalid JSON for DOCX parsing: {e}")
        return {"sections": [], "metadata": {"error": "AI returned invalid JSON"}}
    except Exception as e:
        logger.error(f"DOCX template parsing AI call failed: {e}")
        return {"sections": [], "metadata": {"error": str(e)}}
