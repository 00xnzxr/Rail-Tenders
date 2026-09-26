"""
DRPL Backend - Test Document Service
Handles upload, storage, text extraction, and management of agent test documents.
All documents are persisted on disk + database for reliability.
"""

import os
import logging
from typing import Optional, List
from datetime import datetime, timezone

from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.models.agent_builder import AgentTestDocument

logger = logging.getLogger(__name__)
settings = get_settings()

UPLOAD_SUBDIR = "agent_test_docs"

# Supported MIME types mapped from extensions
MIME_MAP = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".doc": "application/msword",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".tiff": "image/tiff",
    ".tif": "image/tiff",
    ".bmp": "image/bmp",
    ".txt": "text/plain",
    ".md": "text/markdown",
}

SUPPORTED_EXTENSIONS = set(MIME_MAP.keys())


def _get_upload_dir(agent_id: int) -> str:
    """Get (and create) the upload directory for an agent's test documents."""
    upload_dir = os.path.join(settings.upload_dir, UPLOAD_SUBDIR, str(agent_id))
    os.makedirs(upload_dir, exist_ok=True)
    return upload_dir


# ── Text Extraction ──────────────────────────────────────────────────────────


def _extract_pdf_text(file_path: str) -> dict:
    """Extract text from PDF with OCR/vision fallback for scanned documents.

    Cascades through:
      1. pdfplumber for text-searchable pages,
      2. Claude vision (cached per page-sha1) for sparse pages,
      3. Tesseract OCR as a final fallback.

    This means scanned tender PDFs uploaded to the Agent Builder test
    panel (or any consumer of `build_documents_context`) get usable text
    instead of `[No text could be extracted]`, which downstream agents
    treat as a hard "document unreadable" halt.
    """
    try:
        from app.services.advanced_document_parser import (
            extract_text_from_pdf_advanced,
        )
        result = extract_text_from_pdf_advanced(file_path)
        pages = result.get("pages") or []
        page_count = result.get("total_pages", len(pages))

        method_counts: dict[str, int] = {}
        page_chunks: list[str] = []
        for p in pages:
            text = (p.get("text") or "").strip()
            if not text:
                continue
            method = p.get("method") or "text"
            method_counts[method] = method_counts.get(method, 0) + 1
            page_chunks.append(f"[page {p.get('page_num')} · {method}]\n{text}")

        full_text = "\n\n".join(page_chunks) if page_chunks else ""
        # Pick a representative method label: "ocr" if any page used OCR
        # or vision, otherwise "pdfplumber". Used by the UI / status rows.
        if any(m in ("vision", "ocr") for m in method_counts):
            extraction_method = "ocr"
        elif method_counts.get("text"):
            extraction_method = "pdfplumber"
        else:
            extraction_method = "empty"

        if full_text:
            logger.info(
                f"[DRPL] PDF extraction for {file_path}: "
                f"{len(full_text)} chars, methods={method_counts}, pages={page_count}"
            )
        return {
            "text": full_text,
            "page_count": page_count,
            "method": extraction_method,
            "status": "success" if full_text else "partial",
        }
    except Exception as e:
        logger.error(f"[DRPL] PDF extraction failed for {file_path}: {e}")
        return {
            "text": "",
            "page_count": 0,
            "method": "pdfplumber",
            "status": "failed",
            "error": str(e),
        }


def _extract_docx_text(file_path: str) -> dict:
    """Extract text from DOCX using python-docx."""
    try:
        from docx import Document
        doc = Document(file_path)
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

        # Also extract table content
        table_texts = []
        for table in doc.tables:
            rows = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                rows.append(" | ".join(cells))
            table_texts.append("\n".join(rows))

        full_text = "\n".join(paragraphs)
        if table_texts:
            full_text += "\n\n[TABLES]\n" + "\n\n".join(table_texts)

        return {
            "text": full_text,
            "page_count": max(1, len(paragraphs) // 40 + 1),
            "method": "docx",
            "status": "success" if full_text else "partial",
        }
    except ImportError:
        return {"text": "", "page_count": 0, "method": "docx", "status": "failed", "error": "python-docx not installed"}
    except Exception as e:
        logger.error(f"[DRPL] DOCX extraction failed for {file_path}: {e}")
        return {"text": "", "page_count": 0, "method": "docx", "status": "failed", "error": str(e)}


def _extract_image_text(file_path: str) -> dict:
    """Extract text from image using OCR (pytesseract)."""
    try:
        from PIL import Image
        import pytesseract
        image = Image.open(file_path)
        text = pytesseract.image_to_string(image, lang="eng")
        return {
            "text": text.strip() if text else "",
            "page_count": 1,
            "method": "ocr",
            "status": "success" if text and text.strip() else "partial",
        }
    except ImportError:
        return {"text": "", "page_count": 1, "method": "ocr", "status": "failed", "error": "pytesseract/Pillow not installed"}
    except Exception as e:
        logger.error(f"[DRPL] Image OCR failed for {file_path}: {e}")
        return {"text": "", "page_count": 1, "method": "ocr", "status": "failed", "error": str(e)}


def _extract_plain_text(file_path: str) -> dict:
    """Read plain text files."""
    try:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            text = f.read()
        return {
            "text": text,
            "page_count": 1,
            "method": "text",
            "status": "success" if text else "partial",
        }
    except Exception as e:
        return {"text": "", "page_count": 1, "method": "text", "status": "failed", "error": str(e)}


def extract_text(file_path: str) -> dict:
    """
    Route to the correct extractor based on file extension.
    Returns: {"text": str, "page_count": int, "method": str, "status": str, "error"?: str}
    """
    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".pdf":
        return _extract_pdf_text(file_path)
    elif ext in (".docx", ".doc"):
        return _extract_docx_text(file_path)
    elif ext in (".jpg", ".jpeg", ".png", ".tiff", ".tif", ".bmp"):
        return _extract_image_text(file_path)
    elif ext in (".txt", ".text", ".md"):
        return _extract_plain_text(file_path)
    else:
        return {"text": "", "page_count": 0, "method": "unknown", "status": "failed", "error": f"Unsupported format: {ext}"}


# ── CRUD Operations ──────────────────────────────────────────────────────────


def upload_test_document(
    db: Session,
    agent_id: int,
    file_name: str,
    file_data: bytes,
    user_id: Optional[int] = None,
) -> AgentTestDocument:
    """
    Save an uploaded file to disk, extract text, and create a DB record.
    Returns the AgentTestDocument record with extracted text.
    """
    ext = os.path.splitext(file_name)[1].lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {ext}")

    if not file_data or len(file_data) == 0:
        raise ValueError("Empty file uploaded")

    # Save file to disk
    upload_dir = _get_upload_dir(agent_id)
    # Avoid filename collisions by prefixing with timestamp
    ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    safe_name = f"{ts}_{file_name}"
    file_path = os.path.join(upload_dir, safe_name)

    with open(file_path, "wb") as f:
        f.write(file_data)

    # Extract text
    result = extract_text(file_path)

    mime_type = MIME_MAP.get(ext, "application/octet-stream")

    doc = AgentTestDocument(
        agent_id=agent_id,
        file_name=file_name,
        file_path=os.path.abspath(file_path),
        file_size=len(file_data),
        mime_type=mime_type,
        extracted_text=result["text"] if result["text"] else None,
        page_count=result.get("page_count"),
        extraction_method=result.get("method"),
        extraction_status=result.get("status", "failed"),
        extraction_error=result.get("error"),
        uploaded_by=user_id,
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)

    logger.info(f"[DRPL] Test document uploaded: {file_name} (agent={agent_id}, status={result.get('status')})")
    return doc


def list_test_documents(db: Session, agent_id: int) -> List[AgentTestDocument]:
    """List all test documents for an agent, newest first."""
    return (
        db.query(AgentTestDocument)
        .filter(AgentTestDocument.agent_id == agent_id)
        .order_by(AgentTestDocument.uploaded_at.desc())
        .all()
    )


def get_test_document(db: Session, doc_id: int) -> Optional[AgentTestDocument]:
    """Get a single test document by ID."""
    return db.query(AgentTestDocument).filter(AgentTestDocument.id == doc_id).first()


def delete_test_document(db: Session, doc_id: int) -> bool:
    """Delete a test document from disk and database."""
    doc = db.query(AgentTestDocument).filter(AgentTestDocument.id == doc_id).first()
    if not doc:
        return False

    # Remove file from disk
    if doc.file_path and os.path.exists(doc.file_path):
        try:
            os.remove(doc.file_path)
        except OSError as e:
            logger.warning(f"[DRPL] Could not delete file {doc.file_path}: {e}")

    db.delete(doc)
    db.commit()
    return True


def delete_all_test_documents(db: Session, agent_id: int) -> int:
    """Delete all test documents for an agent. Returns count deleted."""
    docs = db.query(AgentTestDocument).filter(AgentTestDocument.agent_id == agent_id).all()
    count = 0
    for doc in docs:
        if doc.file_path and os.path.exists(doc.file_path):
            try:
                os.remove(doc.file_path)
            except OSError:
                pass
        db.delete(doc)
        count += 1
    db.commit()
    return count


def re_extract_text(db: Session, doc_id: int) -> Optional[AgentTestDocument]:
    """Re-run text extraction on a previously uploaded document."""
    doc = db.query(AgentTestDocument).filter(AgentTestDocument.id == doc_id).first()
    if not doc or not doc.file_path or not os.path.exists(doc.file_path):
        return None

    result = extract_text(doc.file_path)
    doc.extracted_text = result["text"] if result["text"] else None
    doc.page_count = result.get("page_count")
    doc.extraction_method = result.get("method")
    doc.extraction_status = result.get("status", "failed")
    doc.extraction_error = result.get("error")
    db.commit()
    db.refresh(doc)
    return doc


def build_documents_context(db: Session, agent_id: int, doc_ids: Optional[List[int]] = None) -> dict:
    """
    Build the document context string for agent execution.
    If doc_ids is provided, only include those documents. Otherwise include all for the agent.
    Returns: {"uploaded_documents": str, "uploaded_file_names": list, "document_count": int}
    """
    if doc_ids:
        docs = db.query(AgentTestDocument).filter(
            AgentTestDocument.id.in_(doc_ids),
            AgentTestDocument.agent_id == agent_id,
        ).all()
    else:
        docs = list_test_documents(db, agent_id)

    documents_text = []
    file_names = []

    for doc in docs:
        if doc.extracted_text:
            documents_text.append(f"--- Document: {doc.file_name} ---\n{doc.extracted_text}")
        elif doc.extraction_status == "failed":
            error_msg = doc.extraction_error or "Unknown extraction error"
            documents_text.append(f"--- Document: {doc.file_name} ---\n[Extraction failed: {error_msg}]")
        else:
            documents_text.append(f"--- Document: {doc.file_name} ---\n[No text could be extracted]")
        file_names.append(doc.file_name)

    return {
        "uploaded_documents": "\n\n".join(documents_text) if documents_text else "",
        "uploaded_file_names": file_names,
        "document_count": len(docs),
    }
