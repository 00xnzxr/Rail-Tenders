"""Task 6 — the NIT-mirror xlsx's "Unit Rate"/"Tender Amount" columns must
read the NIT's OWN numbers (`tender_rate`/`tender_amount`), never the firm's
cost estimate (`rate`/`amount`). The firm estimate gets its own separate
"Est. Rate"/"Est. Amount" columns. Also covers the "Reconciliation" section
sourced from the breakdown's `reconciliation` field (Task 5)."""
import pytest

pytest.importorskip("openpyxl")

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.services.langchain.tools.xlsx_generator_tool import build_cost_xlsx


def _row(tender_rate=100, rate=80, qty=2):
    return {
        "sr_no": 1,
        "item_code": "A-1",
        "schedule_name": "A",
        "description": "Item one",
        "quantity": qty,
        "unit": "no",
        "tender_rate": tender_rate,
        "rate": rate,
        "basic_value": tender_rate * qty,
        "escalation_pct": 0,
        "bidding_unit": "no",
        "is_tax_line": False,
    }


def _header_col_map(ws, header_row=3):
    return {
        ws.cell(row=header_row, column=c).value: c
        for c in range(1, ws.max_column + 1)
        if ws.cell(row=header_row, column=c).value
    }


def test_nit_mirror_single_sheet_tender_columns_read_nit_rate(tmp_path):
    fpath = str(tmp_path / "nit.xlsx")
    rows = [_row(tender_rate=100, rate=80, qty=2)]

    build_cost_xlsx(fpath, "Cost Breakdown", rows, single_sheet=True, layout="nit_mirror")

    wb = load_workbook(fpath, data_only=False)
    ws = wb["Schedule of Items"]

    headers = {}
    for r in range(1, 10):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if v in ("Unit Rate", "Tender Amount", "Est. Rate", "Est. Amount"):
                headers[v] = (r, c)

    assert "Unit Rate" in headers
    assert "Tender Amount" in headers
    assert "Est. Rate" in headers
    assert "Est. Amount" in headers

    unit_rate_row, unit_rate_col = headers["Unit Rate"]
    data_row = unit_rate_row + 1  # first data row right below the header

    # Tender "Unit Rate" must be the NIT rate (100), NOT the firm rate (80).
    assert ws.cell(row=data_row, column=unit_rate_col).value == 100

    tender_amount_col = headers["Tender Amount"][1]
    tender_amount_cell = ws.cell(row=data_row, column=tender_amount_col).value
    # Formula (live, so editing Unit Rate ripples through) referencing the
    # Unit Rate column (NIT tender_rate) — must resolve to qty * tender_rate
    # = 2 * 100 = 200, and must NOT reference the Est. Rate column (firm rate).
    unit_rate_ref = f"{get_column_letter(unit_rate_col)}{data_row}"
    est_rate_col = headers["Est. Rate"][1]
    est_amount_col = headers["Est. Amount"][1]
    est_rate_ref = f"{get_column_letter(est_rate_col)}{data_row}"
    assert unit_rate_ref in str(tender_amount_cell)
    assert est_rate_ref not in str(tender_amount_cell)

    assert ws.cell(row=data_row, column=est_rate_col).value == 80
    est_amount_cell = ws.cell(row=data_row, column=est_amount_col).value
    # Est. Amount must reference the Est. Rate column, not Unit Rate.
    assert est_rate_ref in str(est_amount_cell)
    assert unit_rate_ref not in str(est_amount_cell)


def test_nit_mirror_multi_sheet_tender_columns_read_nit_rate(tmp_path):
    fpath = str(tmp_path / "nit_multi.xlsx")
    rows = [_row(tender_rate=100, rate=80, qty=2)]

    build_cost_xlsx(fpath, "Cost Breakdown", rows, single_sheet=False, layout="nit_mirror")

    wb = load_workbook(fpath, data_only=False)
    ws = wb["Schedule A"]

    header_row = 3
    headers = {
        ws.cell(row=header_row, column=c).value: c
        for c in range(1, ws.max_column + 1)
    }
    assert "Unit Rate" in headers
    assert "Tender Amount" in headers
    assert "Est. Rate" in headers
    assert "Est. Amount" in headers

    data_row = header_row + 1
    assert ws.cell(row=data_row, column=headers["Unit Rate"]).value == 100
    tender_amount_cell = ws.cell(row=data_row, column=headers["Tender Amount"]).value
    unit_rate_ref = f"{get_column_letter(headers['Unit Rate'])}{data_row}"
    est_rate_ref = f"{get_column_letter(headers['Est. Rate'])}{data_row}"
    assert unit_rate_ref in str(tender_amount_cell)
    assert est_rate_ref not in str(tender_amount_cell)
    assert ws.cell(row=data_row, column=headers["Est. Rate"]).value == 80


def test_reconciliation_block_renders_pass_and_flag(tmp_path):
    fpath = str(tmp_path / "nit_reconcile.xlsx")
    rows = [_row(tender_rate=100, rate=80, qty=2)]
    reconciliation = {
        "ok": False,
        "schedules": [
            {"code": "A", "line_sum": 200.0, "stated_total": 200.0, "delta": 0.0, "ok": True},
            {"code": "B", "line_sum": 150.0, "stated_total": 300.0, "delta": -150.0, "ok": False},
        ],
        "grand_line_sum": 350.0,
        "advertised_value": 500.0,
        "grand_delta": -150.0,
    }

    build_cost_xlsx(
        fpath, "Cost Breakdown", rows,
        single_sheet=True, layout="nit_mirror",
        breakdown_meta={"reconciliation": reconciliation},
    )

    wb = load_workbook(fpath, data_only=False)
    ws = wb["Schedule of Items"]

    all_values = [
        ws.cell(row=r, column=c).value
        for r in range(1, ws.max_row + 1)
        for c in range(1, ws.max_column + 1)
    ]
    assert "Reconciliation" in all_values
    assert "PASS" in all_values
    assert "FLAG" in all_values


def test_no_reconciliation_data_renders_no_block(tmp_path):
    fpath = str(tmp_path / "nit_no_reconcile.xlsx")
    rows = [_row(tender_rate=100, rate=80, qty=2)]

    build_cost_xlsx(fpath, "Cost Breakdown", rows, single_sheet=True, layout="nit_mirror")

    wb = load_workbook(fpath, data_only=False)
    ws = wb["Schedule of Items"]
    all_values = [
        ws.cell(row=r, column=c).value
        for r in range(1, ws.max_row + 1)
        for c in range(1, ws.max_column + 1)
    ]
    assert "Reconciliation" not in all_values
