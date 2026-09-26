"""An annexure is the breakdown of the schedule item that cites it.

The Liluah tender, fourth look. The NIT prices "Material Cost for Conversion
work from ICF to NMGHSR coaches (As per Annexure-II of Material list uploaded
in Document Section of NIT)" at Rs 8,16,051 per coach set, fifteen coach
sets. Annexure-II is the thirty materials that make up one coach set.
Costing each material is what the bidder asked for. Adding those thirty
amounts to a total that also holds the item they compose is what the
platform did once it could finally read them: the estimated cost carried the
annexure twice, and the margin view showed the components as a negative
"Other" bucket.

Pinned here, top to bottom:

  - a row read under "Annexure-II" is tagged II, and "2", "ii", "Annnexure- IV"
    (the NIT's own typo) all normalise;
  - the schedule item citing an annexure becomes the parent of its rows, by
    position, so a repeated word code ("CONVERSION MAT" on six rows) cannot
    bind the wrong item;
  - two annexures both starting at serial 1 survive cross-document dedup;
  - the skeleton groups components under their parent and the totals never
    count them;
  - the sum of a parent's priced components (per set) is the parent's rate,
    times its own quantity its amount, and against the published rate its
    margin -- and a user editing a component's rate moves the parent on save;
  - a parent with components is not sent to the agent for its own research;
  - the workbook's grand total leaves component groups out;
  - the capture report names what arrived and what the schedule cites but
    nobody uploaded;
  - the serial-only row branch no longer accepts a compliance-matrix cell,
    and takes serial-only rows only in table context.
"""

import itertools

import pytest

from app.models.cost_breakdown import CostBreakdown, CostBreakdownLine
from app.models.costing_template import BOQItem
from app.models.tender import Tender, TenderDocument
from app.services import boq_parser_service as bps
from app.services import cost_breakdown_service as cbs

pytest.importorskip("openpyxl")

_ids = itertools.count(700_000)


@pytest.fixture
def tid():
    return next(_ids)


# ── labels ───────────────────────────────────────────────────────────────────


@pytest.mark.parametrize("raw,expected", [
    ("2", "II"), ("II", "II"), ("ii", "II"), ("7", "VII"), ("VII", "VII"),
    ("Annexure-VII", "VII"), ("Annnexure- IV", "IV"), ("Annexure - B", "B"),
    ("B", "B"), ("VIB", None), ("0", None), ("", None), (None, None),
])
def test_annexure_labels_normalise(raw, expected):
    assert bps.normalize_annexure_ref(raw) == expected


def test_a_schedule_item_cites_its_annexure():
    d = ("Material Cost for Conversion work from ICF to NMGHSR coaches (As per "
         "Annexure-II of Material list uploaded in Document Section of NIT).")
    assert bps._cited_annexures(d) == ["II"]
    assert bps._cited_annexures("Striping (As per Annnexure- IV of Material list)") == ["IV"]
    assert bps._cited_annexures("Labour Cost of Corrosion work as per Scope of Work.") == []


def test_annexure_headings_carry_across_pages():
    pages = [
        {"text": "EASTERN RAILWAY\nAnnexure-I\nSl Description"},
        {"text": "rows continue"},
        {"text": "tail of I\nANNEXURE - II\nSl No Description"},
        {"text": "rows of II"},
    ]
    carry, starts = bps._annexure_carry(pages)
    assert carry == [None, "I", "I", "II"]
    assert starts == [True, False, True, False]


# ── linking ──────────────────────────────────────────────────────────────────


def _nit_rows():
    return [
        {"sr_no": 1, "item_code": "STRIPPING", "schedule_name": "A", "quantity": 15,
         "unit": "Coach Set", "estimated_rate": 38000,
         "description": "Stripping and Oxy cutting (As per Annexure-I of Material list)"},
        {"sr_no": 2, "item_code": "CORROSION", "schedule_name": "A", "quantity": 15,
         "unit": "Coach Set", "estimated_rate": 58979.2,
         "description": "Labour Cost of Corrosion work as per Scope of Work."},
        {"sr_no": 3, "item_code": "CONVERSION MAT", "schedule_name": "A", "quantity": 15,
         "unit": "Coach Set", "estimated_rate": 816051,
         "description": "Material Cost for Conversion work (As per Annexure-II of Material list)"},
        {"sr_no": 4, "item_code": "CONVERSION MAT", "schedule_name": "A", "quantity": 300,
         "unit": "Metre", "estimated_rate": 9.48,
         "description": "Steel Wire rope 3 mm"},
        {"sr_no": 4, "item_code": "PAINT MATERIAL", "schedule_name": "B", "quantity": 20,
         "unit": "Coach Set", "estimated_rate": 22547.44,
         "description": "Material Cost for Paint Set consisting of 9 items (As per Annexure-VII of Material list)"},
    ]


def _annexure_rows(ref, n, unit="LTR"):
    return [
        {"sr_no": i, "item_code": None, "schedule_name": None, "annexure_ref": ref,
         "quantity": 10 + i, "unit": unit, "description": f"Material {ref}-{i}"}
        for i in range(1, n + 1)
    ]


def test_components_bind_to_the_citing_item_by_position():
    items = _nit_rows() + _annexure_rows("II", 3) + _annexure_rows("VII", 2)
    report = bps._link_annexure_components(items)

    # Annexure-II -> item at index 2 (sr 3), NOT index 3 which shares the
    # word code "CONVERSION MAT".
    assert all(it["_parent_idx"] == 2 for it in items if it.get("annexure_ref") == "II")
    assert all(it["_parent_idx"] == 4 for it in items if it.get("annexure_ref") == "VII")
    assert report["cited"] == ["I", "II", "VII"]
    assert report["captured"] == {"II": 3, "VII": 2}
    assert report["missing"] == ["I"]
    assert report["unlinked"] == []
    assert [(l["annexure"], l["rows"]) for l in report["linked"]] == [("II", 3), ("VII", 2)]


def test_an_annexure_nobody_cites_stays_standalone():
    items = _nit_rows() + _annexure_rows("IX", 2)
    report = bps._link_annexure_components(items)
    assert not any("_parent_idx" in it for it in items)
    assert report["unlinked"] == ["IX"]


def test_two_annexures_both_starting_at_serial_one_are_not_deduped():
    a = _annexure_rows("I", 3)
    b = _annexure_rows("IV", 3)
    # Same serials, near-identical descriptions (the ICF list and the Hybrid
    # list read alike). Without the annexure in the key these collapsed.
    for x, y in zip(a, b):
        y["description"] = x["description"]
    assert len(bps._dedup_boq_rows(a + b)) == 6
    survivors, dropped, _ = bps._collapse_value_redundant_twins(a + b)
    assert dropped == 0 and len(survivors) == 6


# ── persistence: BOQItem rows carry the link, and the report reads them ─────


def _persist(db, tid, items, docs=None):
    """Write BOQItems the way parse_boq_from_tender does, including the
    synthetic code, the link and the source document."""
    bps._assign_annexure_codes(items)
    bps._link_annexure_components(items)
    records = []
    for it in items:
        b = BOQItem(
            tender_id=tid, sr_no=it["sr_no"], description=it["description"],
            quantity=it.get("quantity"), unit=it.get("unit"),
            estimated_rate=it.get("estimated_rate"), item_code=it.get("item_code"),
            basic_value=it.get("basic_value"),
            schedule_name=it.get("schedule_name"), annexure_ref=it.get("annexure_ref"),
            source_document_id=it.get("_source_doc_id"),
        )
        db.add(b)
        records.append(b)
    db.commit()
    for it, b in zip(items, records):
        if it.get("_parent_idx") is not None:
            b.parent_item_id = records[it["_parent_idx"]].id
    db.commit()
    return records


def _tender(db, tid):
    t = Tender(id=tid, portal="ireps", tender_id=f"T-{tid}", title="Liluah", source_url="x")
    db.add(t)
    db.commit()
    return t


def test_capture_report_says_what_arrived_and_what_is_missing(db, tid):
    _tender(db, tid)
    nit = TenderDocument(tender_id=tid, file_name="NIT.pdf", file_path="k1", mime_type="application/pdf")
    anx = TenderDocument(tender_id=tid, file_name="MaterialListPaintAnnexure-VII.pdf", file_path="k2", mime_type="application/pdf")
    db.add_all([nit, anx])
    db.commit()
    items = _nit_rows() + _annexure_rows("VII", 9)
    for it in items:
        it["_source_doc_id"] = anx.id if it.get("annexure_ref") else nit.id
    _persist(db, tid, items)

    report = bps.schedule_capture_report(db, tid)

    by_name = {d["name"]: d for d in report["documents"]}
    assert by_name["NIT.pdf"]["rows"] == 5 and by_name["NIT.pdf"]["schedules"] == ["A", "B"]
    assert by_name["MaterialListPaintAnnexure-VII.pdf"]["annexures"] == {"VII": 9}
    annex = report["annexures"]
    assert annex["cited"] == ["I", "II", "VII"]
    assert annex["captured"] == {"VII": 9}
    assert annex["missing"] == ["I", "II"]
    assert annex["linked"][0]["annexure"] == "VII"
    assert annex["linked"][0]["parent_schedule"] == "B"
    assert annex["linked"][0]["parent_quantity"] == 20


def test_the_reply_names_the_missing_annexures():
    from app.services.langchain.graphs.chat_agent_wrappers import _format_capture_report

    report = {
        "row_count": 14,
        "documents": [{"name": "NIT.pdf", "rows": 5, "schedules": ["A", "B"], "annexures": {}},
                      {"name": "Paint.pdf", "rows": 9, "schedules": [], "annexures": {"VII": 9}}],
        "annexures": {"cited": ["I", "II", "VII"], "captured": {"VII": 9}, "missing": ["I", "II"],
                      "unlinked": [], "linked": [{"annexure": "VII", "rows": 9, "parent_schedule": "B",
                                                  "parent_code": "PAINT MATERIAL", "parent_description": "Paint set",
                                                  "parent_quantity": 20, "parent_unit": "Coach Set"}]},
    }
    text = "\n".join(_format_capture_report(report))
    assert "NIT.pdf" in text and "Paint.pdf" in text
    assert "Annexure I, II" in text and "no uploaded document" in text
    assert "rolled into Schedule B item PAINT MATERIAL" in text
    assert _format_capture_report({}) == []


# ── the skeleton and the totals ──────────────────────────────────────────────


def _skeleton(db, tid):
    _tender(db, tid)
    items = _nit_rows() + _annexure_rows("VII", 3)
    records = _persist(db, tid, items)
    bd = cbs.build_skeleton_from_boq(db, tid)
    return bd, records


def test_skeleton_groups_components_under_their_parent(db, tid):
    bd, records = _skeleton(db, tid)
    parent = next(r for r in records if r.item_code == "PAINT MATERIAL")
    comps = [ln for ln in bd.lines if ln.parent_boq_item_id == parent.id]
    assert len(comps) == 3
    assert all(ln.annexure_ref == "VII" for ln in comps)
    assert all("components of Schedule B item PAINT MATERIAL" in ln.schedule_section for ln in comps)
    assert all(ln.tender_rate is None for ln in comps), "an annexure prints no published rate"
    assert all(ln.item_code.startswith("ANX-VII-") for ln in comps)
    # The parent is a normal schedule line.
    pl = next(ln for ln in bd.lines if ln.boq_item_id == parent.id)
    assert pl.parent_boq_item_id is None and pl.schedule_section == "Schedule B"


def test_components_are_never_in_the_totals():
    lines = [
        {"boq_item_id": 1, "quantity": 20, "rate": 1000, "amount": 20000, "rate_source": "web_search"},
        {"boq_item_id": 2, "parent_boq_item_id": 1, "quantity": 12, "rate": 50, "amount": 600, "rate_source": "web_search"},
        {"boq_item_id": 3, "parent_boq_item_id": 1, "quantity": 4, "rate": None, "rate_source": "needs_user_input"},
    ]
    totals = cbs._compute_totals(lines, 0, 0, 0)
    assert totals["subtotal"] == 20000
    assert totals["needs_input_count"] == 0, "an unpriced component is the parent's note, not a needs-input line"


def test_rollup_makes_the_parent_rate_the_sum_of_its_components():
    parent = {"boq_item_id": 9, "quantity": 20, "unit": "Coach Set", "tender_rate": 22547.44,
              "escalation_pct": 0, "rate": None, "rate_source": "needs_user_input",
              "schedule_name": "B", "item_code": "PAINT MATERIAL", "sr_no": 4}
    comps = [
        {"boq_item_id": 10, "parent_boq_item_id": 9, "annexure_ref": "VII", "quantity": 12, "rate": 400,
         "amount": 4800, "rate_source": "web_search", "confidence": "high"},
        {"boq_item_id": 11, "parent_boq_item_id": 9, "annexure_ref": "VII", "quantity": 10, "rate": 300,
         "amount": 3000, "rate_source": "training_data", "confidence": "medium"},
    ]
    out = cbs.rollup_component_parents([parent] + comps)

    assert out["parents"] == 1 and out["unpriced_components"] == 0
    assert parent["rate"] == 7800.0, "per set"
    assert parent["amount"] == 156000.0, "20 coach sets"
    assert parent["rate_source"] == cbs.COMPONENT_BUILDUP_SOURCE
    assert parent["needs_input"] is False
    assert parent["confidence"] == "medium", "the weakest component"
    assert parent["tender_amount"] == pytest.approx(450948.8)
    assert parent["margin_amount"] == pytest.approx(450948.8 - 156000.0)
    assert "2 of 2 component" in parent["cost_buildup_note"]


def test_rollup_with_an_unpriced_component_is_a_floor_and_says_so():
    parent = {"boq_item_id": 9, "quantity": 20, "unit": "Coach Set", "tender_rate": 22547.44,
              "rate": None, "rate_source": "needs_user_input", "schedule_name": "B",
              "item_code": "PAINT MATERIAL", "sr_no": 4}
    comps = [
        {"boq_item_id": 10, "parent_boq_item_id": 9, "annexure_ref": "VII", "quantity": 12,
         "rate": 400, "amount": 4800, "rate_source": "web_search", "confidence": "high"},
        {"boq_item_id": 11, "parent_boq_item_id": 9, "annexure_ref": "VII", "quantity": 10,
         "rate": None, "amount": None, "rate_source": "needs_user_input"},
    ]
    out = cbs.rollup_component_parents([parent] + comps)
    assert parent["rate"] == 4800.0 and parent["confidence"] == "low"
    assert "floor" in parent["cost_buildup_note"]
    assert out["notes"] and "1 of 2" in out["notes"][0]


def test_rollup_leaves_a_parent_with_no_priced_component_alone():
    parent = {"boq_item_id": 9, "quantity": 20, "rate": None, "rate_source": "needs_user_input"}
    comp = {"boq_item_id": 10, "parent_boq_item_id": 9, "quantity": 12, "rate": None,
            "rate_source": "needs_user_input"}
    out = cbs.rollup_component_parents([parent, comp])
    assert out["parents"] == 0 and parent["rate"] is None


def test_rollup_persists_and_the_totals_see_only_the_parent(db, tid):
    bd, records = _skeleton(db, tid)
    parent = next(r for r in records if r.item_code == "PAINT MATERIAL")
    # Price the three components the way merge_batch_rates would.
    batch = [
        {"boq_item_id": r.id, "rate": 100.0 * i, "amount": None, "rate_source": "web_search"}
        for i, r in enumerate((r for r in records if r.parent_item_id == parent.id), start=1)
    ]
    cbs.merge_batch_rates(db, bd.id, batch)
    cbs.recompute_breakdown_totals(db, bd)
    db.refresh(bd)

    pl = next(ln for ln in bd.lines if ln.boq_item_id == parent.id)
    comps = [ln for ln in bd.lines if ln.parent_boq_item_id == parent.id]
    per_set = sum(ln.amount for ln in comps)          # 11*100 + 12*200 + 13*300
    assert per_set == pytest.approx(7400.0)
    assert pl.rate == pytest.approx(per_set)
    assert pl.amount == pytest.approx(per_set * 20)
    assert pl.rate_source == cbs.COMPONENT_BUILDUP_SOURCE
    # Only the parent's amount is in the subtotal; the other schedule rows are
    # still unpriced, so the subtotal IS the parent's amount.
    assert bd.subtotal == pytest.approx(per_set * 20)


def test_editing_a_component_moves_the_parent_on_save(db, tid):
    bd, records = _skeleton(db, tid)
    parent = next(r for r in records if r.item_code == "PAINT MATERIAL")
    lines = cbs.to_dict(bd)["lines"]
    for ln in lines:
        if ln["parent_boq_item_id"] == parent.id:
            ln["rate"] = 50.0
            ln["amount"] = None
            ln["needs_input"] = False
            ln["rate_source"] = "user_override"
    updated = cbs.replace_lines(db, bd.id, lines)

    pl = next(ln for ln in updated.lines if ln.boq_item_id == parent.id)
    comps = [ln for ln in updated.lines if ln.parent_boq_item_id == parent.id]
    assert all(ln.parent_boq_item_id == parent.id for ln in comps), "the link survives a save"
    assert pl.rate == pytest.approx(sum(ln.quantity * 50.0 for ln in comps))
    assert updated.subtotal == pytest.approx(pl.amount)


def test_the_edit_payload_carries_the_schedule_link():
    from app.api.routes.cost_breakdown import CostBreakdownLinePayload

    p = CostBreakdownLinePayload(
        description="x", boq_item_id=5, schedule_name="A", item_code="A1",
        is_tax_line=False, parent_boq_item_id=3, annexure_ref="II", source_url="http://x",
    ).model_dump()
    assert p["boq_item_id"] == 5 and p["parent_boq_item_id"] == 3 and p["source_url"] == "http://x"


def test_strategic_summary_has_no_other_bucket_for_components(db, tid):
    bd, records = _skeleton(db, tid)
    summary = cbs.build_strategic_summary(db, bd, tender_id=tid)
    names = [s["schedule"] for s in summary["schedule_breakdown"]]
    assert "Other" not in names
    assert names == ["Schedule A", "Schedule B"]


# ── the agent is not asked to price a parent that is built up ────────────────


def test_parents_with_components_are_not_researched_separately():
    from app.services.langchain.graphs.enhanced_costing_agent import (
        split_costable_rows, _annotate_component_rows, _render_bidding_schedule_block,
    )
    rows = [
        {"boq_item_id": 1, "schedule_name": "B", "item_code": "PAINT MATERIAL", "sr_no": 4,
         "description": "Paint set of 9 items", "quantity": 20, "unit": "Coach Set", "is_tax_line": False},
        {"boq_item_id": 2, "schedule_name": "B", "item_code": "ELECTRICAL LAB", "sr_no": 5,
         "description": "Labour", "quantity": 20, "unit": "Coach Set", "is_tax_line": False},
        {"boq_item_id": 3, "schedule_name": None, "annexure_ref": "VII", "parent_item_id": 1, "sr_no": 1,
         "description": "PU paint", "quantity": 12, "unit": "LTR", "is_tax_line": False},
    ]
    cost_rows, note = split_costable_rows(rows)
    assert note is None
    assert [r["boq_item_id"] for r in cost_rows] == [2, 3]

    _annotate_component_rows(rows)
    assert rows[2]["component_of"].startswith("Schedule B item PAINT MATERIAL")
    assert rows[2]["quantity_basis"] == "per Coach Set"
    block = _render_bidding_schedule_block([rows[2]])
    assert "Annexure-VII components" in block and "PER SET" in block


# ── the workbook ─────────────────────────────────────────────────────────────


def test_workbook_grand_total_leaves_component_groups_out(tmp_path):
    from openpyxl import load_workbook
    from app.services.langchain.tools.xlsx_generator_tool import build_cost_xlsx

    rows = [
        {"sr_no": 4, "description": "Paint set", "qty": 20, "unit": "Coach Set", "tender_rate": 22547.44,
         "tender_amount": 450948.8, "rate": 7800, "amount": 156000, "schedule_section": "Schedule B",
         "boq_item_id": 1, "schedule_name": "B", "item_code": "PAINT MATERIAL"},
        {"sr_no": 1, "description": "PU paint", "qty": 12, "unit": "LTR", "rate": 400, "amount": 4800,
         "schedule_section": "Annexure-VII — components of Schedule B item PAINT MATERIAL (per set; rolled into that item)",
         "boq_item_id": 2, "parent_boq_item_id": 1, "is_component": True, "annexure_ref": "VII", "item_code": "ANX-VII-1"},
    ]
    path = str(tmp_path / "c.xlsx")
    summary = build_cost_xlsx(path, "Liluah", rows, strategic_summary={"tender_snapshot": {}},
                              breakdown_meta={"gst_percent": 18}, single_sheet=True)
    assert summary["total_amount"] == 156000
    assert summary["total_tender_amount"] == pytest.approx(450948.8)

    ws = load_workbook(path)["3. All Schedules"]
    grand = next(c for row in ws.iter_rows() for c in row if c.value == "GRAND TOTAL")
    total_cells = [c for row in ws.iter_rows() for c in row if c.value == "TOTAL"]
    assert len(total_cells) == 2, "each group still shows its own TOTAL"
    # The grand-total formula references exactly one TOTAL row: the schedule's.
    formula = next(
        c.value for c in ws[grand.row] if isinstance(c.value, str) and c.value.startswith("=SUM(")
    )
    assert formula.count(",") == 0, formula


# ── the serial-only branch, hardened ─────────────────────────────────────────


def _serial_row(desc):
    return {"sr_no": 3, "item_code": None, "description": desc, "quantity": None,
            "unit": None, "unit_rate": None, "basic_value": None, "amount": None,
            "schedule_name": None}


@pytest.mark.parametrize("cell", ["Yes No", "No No Not Allowed", "N/A", "Yes", "Not Allowed", "-"])
def test_a_compliance_matrix_cell_is_not_an_item(cell):
    assert not bps._is_valid_schedule_item(_serial_row(cell))


def test_serial_only_rows_need_table_context():
    row = _serial_row("Zinc phosphate primer, IS 2074")
    assert bps._is_valid_schedule_item(row)
    assert not bps._is_valid_schedule_item(row, table_context=False)
    # A quantified row does not depend on context.
    assert bps._is_valid_schedule_item({**row, "quantity": 10, "unit": "LTR"}, table_context=False)


@pytest.mark.asyncio
async def test_chunks_break_at_annexure_headings_and_rows_are_tagged(monkeypatch):
    """Two annexures in one document, two pages each: the walker must cut
    the chunk at the second heading and tag each row with its own annexure,
    without needing the extractor to say which is which."""
    seen: list[str] = []

    async def fake_call(_db, chunk_pages, _carry):
        text = "\n".join(p["text"] for p in chunk_pages)
        seen.append(text)
        return [{"sr_no": i, "description": f"Material {i}", "quantity": i, "unit": "Nos"}
                for i in (1, 2)]

    monkeypatch.setattr(bps, "_call_boq_chunk", fake_call)
    pages = [
        {"page_num": 1, "text": "Annexure-I\nSl No Description Qty Unit\n1 Material 1 1 Nos"},
        {"page_num": 2, "text": "2 Material 2 2 Nos"},
        {"page_num": 3, "text": "Annexure-II\nSl No Description Qty Unit\n1 Material 1 1 Nos"},
        {"page_num": 4, "text": "2 Material 2 2 Nos"},
    ]
    rows = await bps._walk_pages_chunked(
        None, pages, [None] * 4, [0, 1, 2, 3], pages_per_chunk=4, max_retries=0,
    )
    assert len(seen) == 2, "the chunk was cut at the Annexure-II heading"
    assert [r["annexure_ref"] for r in rows] == ["I", "I", "II", "II"]
    assert all("annexure" not in r for r in rows)


@pytest.mark.asyncio
async def test_an_extractor_may_not_invent_an_annexure(monkeypatch):
    async def fake_call(_db, _pages, _carry):
        return [{"sr_no": 1, "description": "Material", "quantity": 1, "unit": "Nos", "annexure": "IX"}]

    monkeypatch.setattr(bps, "_call_boq_chunk", fake_call)
    pages = [{"page_num": 1, "text": "Annexure-VII\nSl No Description Qty Unit\n1 Material 1 Nos"}]
    rows = await bps._walk_pages_chunked(None, pages, [None], [0], pages_per_chunk=4, max_retries=0)
    assert rows[0]["annexure_ref"] == "VII"


@pytest.mark.asyncio
async def test_serial_only_rows_are_refused_inside_an_ireps_schedule_chunk(monkeypatch):
    async def fake_call(_db, _pages, _carry):
        return [
            {"sr_no": 1, "item_code": "A1", "description": "Real item", "quantity": 4, "unit": "Nos"},
            {"sr_no": 7, "description": "No No Not Allowed"},
            {"sr_no": 8, "description": "Some fragment of a merged cell"},
        ]

    monkeypatch.setattr(bps, "_call_boq_chunk", fake_call)
    pages = [{"page_num": 1, "text": "Schedule () A-Engine removal\nItem Code A1 Description:- Real item Item Qty 4"}]
    rows = await bps._walk_pages_chunked(None, pages, ["A"], [0], pages_per_chunk=4, max_retries=0)
    assert [r["sr_no"] for r in rows] == [1]


# ── the re-parse keeps the link ──────────────────────────────────────────────


def test_rebind_after_reparse_keeps_the_component_link(db, tid):
    bd, records = _skeleton(db, tid)
    parent = next(r for r in records if r.item_code == "PAINT MATERIAL")
    pre = bps._snapshot_cost_line_bindings(db, tid)
    assert pre

    # Re-parse: delete and re-insert the same rows with new ids. SQLite
    # recycles ids after a delete (Postgres never does), so push the
    # sequence along with a row on another tender first -- otherwise the
    # "fresh" rows land on the old ids and the rebind is not exercised.
    db.query(BOQItem).filter(BOQItem.tender_id == tid).delete()
    db.commit()
    db.expire_all()
    db.add(BOQItem(tender_id=tid + 500_000, sr_no=1, description="id shifter"))
    db.commit()
    fresh = _persist(db, tid, _nit_rows() + _annexure_rows("VII", 3))
    bps._rebind_cost_lines(db, tid, fresh, pre)
    db.refresh(bd)

    new_parent = next(r for r in fresh if r.item_code == "PAINT MATERIAL")
    assert new_parent.id != parent.id
    comps = [ln for ln in bd.lines if ln.annexure_ref == "VII"]
    assert len(comps) == 3
    assert all(ln.boq_item_id is not None for ln in comps), "annexure rows rebind by their synthetic code"
    assert all(ln.parent_boq_item_id == new_parent.id for ln in comps)


# ── an annexure chunk's sub-table numbers are not schedules ─────────────────


@pytest.mark.asyncio
async def test_a_sub_table_number_in_an_annexure_chunk_is_not_a_schedule(monkeypatch):
    """Annexure-II is thirty numbered sub-tables ("24 | Door Frame Complete").
    The extractor returned the sub-table number as `schedule_name`; the rows
    were persisted as "Schedule 10" / "Schedule 11", never bound to the item
    citing the annexure, and costed as scope of their own."""
    async def fake_call(_db, _pages, _carry):
        return [
            {"sr_no": 1, "schedule_name": "24", "description": "Door Hinge straight-1",
             "quantity": 6, "unit": "kg"},
            {"sr_no": 2, "schedule_name": "24", "description": "Hook Lock", "quantity": 2, "unit": "kg"},
        ]

    monkeypatch.setattr(bps, "_call_boq_chunk", fake_call)
    pages = [{"page_num": 1, "text": "ANNEXURE-II\nS.No Description Qty Unit Weight Rate Total\n24 Door Frame Complete"}]
    rows = await bps._walk_pages_chunked(None, pages, [None], [0], pages_per_chunk=4, max_retries=0)
    assert [(r.get("schedule_name"), r["annexure_ref"]) for r in rows] == [(None, "II"), (None, "II")]


@pytest.mark.asyncio
async def test_a_schedule_name_inside_an_ireps_chunk_is_kept(monkeypatch):
    async def fake_call(_db, _pages, _carry):
        return [{"sr_no": 1, "schedule_name": "B", "item_code": "B1", "description": "Real",
                 "quantity": 4, "unit": "Nos"}]

    monkeypatch.setattr(bps, "_call_boq_chunk", fake_call)
    pages = [{"page_num": 1, "text": "Schedule () B-Conversion\nItem Code B1 Description:- Real Item Qty 4\nAnnexure-IV"}]
    rows = await bps._walk_pages_chunked(None, pages, [None], [0], pages_per_chunk=4, max_retries=0)
    assert rows[0]["schedule_name"] == "B" and "annexure_ref" not in rows[0]


def test_an_annexure_row_after_the_schedule_is_not_filled_with_the_running_schedule():
    items = [
        {"schedule_name": "B", "description": "Paint set (Annexure-VII)"},
        {"schedule_name": None, "annexure_ref": "VII", "description": "Primer"},
        {"schedule_name": None, "description": "continuation of B"},
    ]
    bps._normalize_schedule_names(items)
    assert [it.get("schedule_name") for it in items] == ["B", None, "B"]


# ── the order Rajesh reads it in: schedules, then Annexure-I, II, ... ────────


def test_display_order_is_schedules_then_annexures_in_numeral_order():
    """The ORM orders lines by serial alone, which put a component with
    serial 1 above Schedule A and the groups on the sheet as "Schedule 11,
    B, A, 10". Wanted: A, B, then Annexure-I, II, ... VII, each annexure in
    the order its table was read (Annexure-II restarts its serial per
    sub-assembly, so a serial sort would interleave thirty tables)."""
    lines = [
        {"sr_no": 1, "schedule_name": "", "parent_boq_item_id": 7, "annexure_ref": "VII", "boq_item_id": 40},
        {"sr_no": 1, "schedule_name": "B", "boq_item_id": 12},
        {"sr_no": 2, "schedule_name": "A", "boq_item_id": 2},
        {"sr_no": 1, "schedule_name": "", "parent_boq_item_id": 3, "annexure_ref": "II", "boq_item_id": 21},
        {"sr_no": 1, "schedule_name": "A", "boq_item_id": 1},
        {"sr_no": 1, "schedule_name": "", "parent_boq_item_id": 3, "annexure_ref": "II", "boq_item_id": 25},
        {"sr_no": 3, "schedule_name": "", "parent_boq_item_id": 3, "annexure_ref": "II", "boq_item_id": 23},
        {"sr_no": 2, "schedule_name": "", "parent_boq_item_id": 1, "annexure_ref": "I", "boq_item_id": 18},
        {"sr_no": 10, "schedule_name": "A", "boq_item_id": 10},
        {"sr_no": 5, "schedule_name": "", "annexure_ref": "IX", "boq_item_id": 60},
        {"sr_no": 1, "schedule_name": "", "boq_item_id": None, "description": "stray"},
        {"sr_no": 1, "schedule_name": "A7", "boq_item_id": 30},
    ]
    out = cbs.order_lines_for_display(lines)
    assert [l.get("boq_item_id") for l in out] == [1, 2, 10, 30, 12, 18, 21, 23, 25, 40, 60, None]


def test_display_order_accepts_orm_rows(db, tid):
    t = _tender(db, tid)
    bd = CostBreakdown(tender_id=t.id, title="x", created_by_agent="t")
    db.add(bd); db.flush()
    db.add_all([
        CostBreakdownLine(cost_breakdown_id=bd.id, sr_no=1, description="component", quantity=1,
                          parent_boq_item_id=1, annexure_ref="II", boq_item_id=99),
        CostBreakdownLine(cost_breakdown_id=bd.id, sr_no=2, description="A2", quantity=1,
                          schedule_name="A", boq_item_id=2),
        CostBreakdownLine(cost_breakdown_id=bd.id, sr_no=1, description="A1", quantity=1,
                          schedule_name="A", boq_item_id=1),
    ])
    db.flush(); db.expire_all()
    bd = db.get(CostBreakdown, bd.id)
    assert [l.description for l in cbs.order_lines_for_display(bd.lines)] == ["A1", "A2", "component"]
    assert [l["description"] for l in cbs.to_dict(bd)["lines"]] == ["A1", "A2", "component"]


# ── Annexure-II restarts its serial in every sub-table ──────────────────────


def test_annexure_codes_number_the_rows_in_reading_order():
    """Thirty sub-tables each start at serial 1; the code must not."""
    items = [
        {"sr_no": 1, "annexure_ref": "II", "description": "Channel"},
        {"sr_no": 2, "annexure_ref": "II", "description": "Hook"},
        {"sr_no": 1, "annexure_ref": "II", "description": "Door Hinge straight-1"},
        {"sr_no": 1, "annexure_ref": "VII", "description": "Primer"},
        {"sr_no": 1, "schedule_name": "A", "item_code": "STRIPPING", "description": "nit row"},
    ]
    bps._assign_annexure_codes(items)
    assert [it.get("item_code") for it in items] == [
        "ANX-II-1", "ANX-II-2", "ANX-II-3", "ANX-VII-1", "STRIPPING",
    ]


def test_dedup_keeps_a_recurring_part_name_across_sub_tables():
    rows = [
        {"sr_no": 2, "annexure_ref": "II", "description": "Hook Lock", "quantity": 2, "basic_value": 5.63},
        {"sr_no": 2, "annexure_ref": "II", "description": "Hook Lock", "quantity": 4, "basic_value": 11.26},
        {"sr_no": 2, "annexure_ref": "II", "description": "Hook Lock", "quantity": 2, "basic_value": 5.63},
    ]
    out = bps._dedup_boq_rows(rows)
    assert [r["quantity"] for r in out] == [2, 4], "the third is the first again; the second is another sub-table's"


def test_twin_collapse_does_not_merge_different_parts_that_share_a_serial():
    """(ANX:II, 1) holds one row per sub-table. None is a twin of another,
    none is a conflict to report, and a title row (no value) is absorbed only
    by a valued row of the same description."""
    rows = [
        {"sr_no": 1, "annexure_ref": "II", "description": "Channel for middle trough", "quantity": 2, "basic_value": 460.21},
        {"sr_no": 1, "annexure_ref": "II", "description": "Door Hinge straight-1", "quantity": 6, "basic_value": 7535.5},
        {"sr_no": 1, "annexure_ref": "II", "description": "Plain washer", "quantity": 2, "basic_value": 460.21},
        {"sr_no": 1, "annexure_ref": "II", "description": "Door Hinge straight-1", "quantity": 6, "basic_value": 7535.5},
        {"sr_no": 1, "annexure_ref": "II", "description": "Door Hinge straight-1"},
        {"sr_no": 1, "annexure_ref": "II", "description": "Door Frame Complete"},
    ]
    out, dropped, unresolvable = bps._collapse_value_redundant_twins(rows)
    assert unresolvable == []
    assert dropped == 2
    assert [r["description"] for r in out] == [
        "Channel for middle trough", "Door Hinge straight-1", "Plain washer", "Door Frame Complete",
    ]


def test_twin_collapse_inside_a_schedule_is_unchanged():
    rows = [
        {"sr_no": 3, "schedule_name": "A", "item_code": "A3", "description": "x", "quantity": 1, "basic_value": 10.0},
        {"sr_no": 3, "schedule_name": "A", "description": "x", "quantity": 1},
        {"sr_no": 4, "schedule_name": "A", "item_code": "A4", "description": "y", "quantity": 1, "basic_value": 10.0},
        {"sr_no": 4, "schedule_name": "A", "item_code": "A4b", "description": "z", "quantity": 1, "basic_value": 20.0},
    ]
    out, dropped, unresolvable = bps._collapse_value_redundant_twins(rows)
    assert dropped == 1 and [g["sr_no"] for g in unresolvable] == [4]


def test_a_sub_table_number_in_the_code_column_is_replaced_by_the_sequence_code():
    items = [
        {"sr_no": 1, "annexure_ref": "II", "item_code": "12", "description": "Bottom angle"},
        {"sr_no": 2, "annexure_ref": "II", "item_code": "12", "description": "Guide"},
        {"sr_no": 1, "annexure_ref": "II", "item_code": "MAT-7", "description": "keeps a real code"},
    ]
    bps._assign_annexure_codes(items)
    assert [it["item_code"] for it in items] == ["ANX-II-1", "ANX-II-2", "MAT-7"]


@pytest.mark.asyncio
async def test_rows_before_a_mid_chunk_heading_keep_the_earlier_annexure(monkeypatch):
    """The tail of Annexure-II's last table shares a page with the ANNEXURE-III
    heading. Unlabelled rows take the annexure the chunk is in at that point
    in document order, not the last heading on the page."""
    async def fake_call(_db, chunk_pages, _carry):
        text = "\n".join(p["text"] for p in chunk_pages)
        rows = [
            {"sr_no": 1, "description": "LP sheet"},
            {"sr_no": 11, "description": "Bib Cock"},
            {"sr_no": 12, "description": "Carton sheet 2mm"},
            {"sr_no": 1, "description": "Complete Pull Rod for Hand Brake", "quantity": 1, "unit": "Set", "annexure": "III"},
            {"sr_no": 2, "description": "Hand Brake arrangement", "quantity": 1, "unit": "Set"},
        ]
        return [r for r in rows if r["description"].split()[0] in text]

    monkeypatch.setattr(bps, "_call_boq_chunk", fake_call)
    pages = [
        {"page_num": 12, "text": "ANNEXURE-II\nS.No Description Qty Unit\n1 LP sheet"},
        {"page_num": 13, "text": "11 Bib Cock\n12 Carton sheet 2mm\nANNEXURE-III\nS.No Description Qty Unit\n1 Complete Pull Rod\n2 Hand Brake arrangement"},
    ]
    rows = await bps._walk_pages_chunked(None, pages, [None, None], [0, 1], pages_per_chunk=4, max_retries=0)
    assert [r["annexure_ref"] for r in rows] == ["II", "II", "II", "III", "III"]


@pytest.mark.parametrize("desc", [
    "Louver-1 RDSO/CG/DRG/21032",
    "Fall plate arrangement RDSO/CG/DRG/20006, alt.4",
    "Barrel Locking arrangement RDSO/CG/DRG/20010",
    "Hand brake lever as per Drg. No. K-1034",
])
def test_a_drawing_register_row_is_not_an_item(desc):
    assert not bps._is_valid_schedule_item({"sr_no": 27, "description": desc}, table_context=True)


def test_a_material_row_citing_its_drawing_is_still_an_item():
    assert bps._is_valid_schedule_item(
        {"sr_no": 3, "description": "Louver frame arrangement as per Drg. No. RDSO/CG/DRG/20008", "quantity": 8, "unit": "Nos"},
        table_context=True,
    )
    assert bps._is_valid_schedule_item(
        {"sr_no": 1, "description": "Complete Pull Rod for Hand Brake as per RDSO Drg.No Sketch K-1034", "unit": "Set"},
        table_context=True,
    )


def test_capture_report_says_how_much_of_the_annexure_value_was_read(db, tid):
    """Annexure-II prints a value per row; item A3 prints Rs 8,16,051 per set.
    Their ratio is the completeness signal a row count cannot give."""
    t = _tender(db, tid)
    items = _nit_rows()
    comps = [
        {"sr_no": 1, "description": "Channel", "quantity": 15.712, "unit": "kg",
         "basic_value": 460.21, "annexure_ref": "II"},
        {"sr_no": 2, "description": "Longitudinal Channel", "quantity": 75.25, "unit": "kg",
         "basic_value": 407565.29, "annexure_ref": "II"},
        {"sr_no": 3, "description": "no printed value", "annexure_ref": "II"},
    ]
    _persist(db, t.id, items + comps)
    rep = bps.schedule_capture_report(db, t.id)
    (link,) = [l for l in rep["annexures"]["linked"] if l["annexure"] == "II"]
    assert link["rows"] == 3
    assert link["published_rate"] == 816051.0
    assert link["captured_value"] == 408025.5
    assert link["coverage_pct"] == 50.0
