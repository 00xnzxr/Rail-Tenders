"""The costing Excel must not contain margin or overhead columns/roll-ups.

Client requirement: the downloadable costing sheet should show line costs, GST,
and Total only — no Margin / Overhead columns and no Overhead/Margin footer
roll-ups. Tender Rate / Tender Amount stay in the margin-analysis sheet as
reference; the Summary sheet's strategic profitability table is out of scope.
"""

import pytest

pytest.importorskip("openpyxl")
from openpyxl import load_workbook

from app.services.langchain.tools.xlsx_generator_tool import (
    build_cost_xlsx,
    _COST_COLUMNS,
    _MARGIN_COLUMNS,
)

_MARGIN_KEYS = {"margin_amount", "margin_pct", "profit_pct", "profit_amount"}


def _keys(cols):
    return {k for k, _, _ in cols}


def _string_cells(path, skip_prefixes=()):
    wb = load_workbook(path)
    out = []
    for ws in wb.worksheets:
        if any(ws.title.startswith(p) for p in skip_prefixes):
            continue
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if isinstance(v, str):
                    out.append(v)
    return out


def test_cost_columns_drop_margin_keys():
    assert _keys(_COST_COLUMNS).isdisjoint(_MARGIN_KEYS)


def test_margin_columns_drop_margin_keys_but_keep_tender_and_cost():
    ks = _keys(_MARGIN_COLUMNS)
    assert ks.isdisjoint({"margin_amount", "margin_pct", "profit_pct"})
    # Tender Rate/Amount and Est. cost columns stay.
    assert {"tender_rate", "tender_amount", "rate", "amount"} <= ks


def test_single_sheet_build_has_no_margin_or_overhead(tmp_path):
    path = str(tmp_path / "single.xlsx")
    build_cost_xlsx(path, "Cost Breakdown", [
        {"sr_no": 1, "description": "Excavation", "qty": 10, "unit": "cum", "rate": 250, "amount": 2500},
        {"sr_no": 2, "description": "Concrete", "qty": 5, "unit": "cum", "rate": 6000, "amount": 30000},
    ])
    joined = " | ".join(_string_cells(path)).lower()
    assert "margin" not in joined
    assert "overhead" not in joined
    # Still a working cost sheet.
    assert "gst" in joined
    assert "subtotal" in joined or "grand total" in joined


def test_margin_analysis_schedule_sheet_has_no_margin_columns(tmp_path):
    path = str(tmp_path / "margin.xlsx")
    # tender_rate (without item_code+schedule_name) routes to the margin-analysis
    # multi-sheet builder — the layout the client actually downloads.
    rows = [
        {"sr_no": 1, "description": "Manpower", "unit": "no", "qty": 4,
         "tender_rate": 150, "tender_amount": 600, "rate": 100, "amount": 400},
        {"sr_no": 2, "description": "Material", "unit": "kg", "qty": 10,
         "tender_rate": 90, "tender_amount": 900, "rate": 70, "amount": 700},
    ]
    build_cost_xlsx(path, "Cost Breakdown", rows)
    # Skip the Summary sheet (strategic gross-margin table is out of scope).
    joined = " | ".join(_string_cells(path, skip_prefixes=("1.",))).lower()
    assert "margin" not in joined, f"margin label leaked: {joined}"
    assert "profit %" not in joined
    # Cost + tender reference columns remain.
    assert "est. total cost" in joined or "est total cost" in joined
    assert "tender rate" in joined
