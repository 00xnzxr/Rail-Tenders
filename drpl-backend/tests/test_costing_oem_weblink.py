"""Piece 3 — OEM manufacturer + web link as first-class costing line fields."""
from app.models.cost_breakdown import CostBreakdownLine
from app.services.cost_breakdown_service import _normalize_line_dict


def test_model_has_oem_and_source_url_columns():
    cols = {c.name for c in CostBreakdownLine.__table__.columns}
    assert "oem_manufacturer" in cols
    assert "source_url" in cols


def test_normalize_line_dict_carries_oem_and_source_url():
    line = {
        "description": "LED fitting", "quantity": 10, "rate": 500,
        "rate_source": "web_search",
        "oem_manufacturer": "Wipro Lighting",
        "source_url": "https://example.com/led-fitting",
    }
    # _normalize_line_dict(line, sr_no, breakdown_margin_percent)
    out = _normalize_line_dict(line, 1, 0.0)
    assert out["oem_manufacturer"] == "Wipro Lighting"
    assert out["source_url"] == "https://example.com/led-fitting"


def test_normalize_line_dict_defaults_missing_to_none():
    out = _normalize_line_dict({"description": "X", "quantity": 1, "rate": 1}, 1, 0.0)
    assert out["oem_manufacturer"] is None
    assert out["source_url"] is None


import openpyxl
from app.services.langchain.tools.xlsx_generator_tool import build_cost_xlsx


def test_single_sheet_has_oem_and_clickable_web_source(tmp_path):
    out = tmp_path / "c.xlsx"
    build_cost_xlsx(
        str(out), "Cost",
        rows=[{"sr_no": "1", "description": "LED", "qty": 2, "rate": 500,
               "amount": 1000, "rate_source": "web_search",
               "oem_manufacturer": "Wipro",
               "source_url": "https://example.com/led"}],
    )
    wb = openpyxl.load_workbook(str(out))
    ws = wb.active
    header = [c.value for row in ws.iter_rows(min_row=1, max_row=ws.max_row) for c in row]
    assert "OEM / Manufacturer" in header
    assert "Web Source" in header
    # OEM text present; Web Source cell is a hyperlink.
    texts = [c.value for r in ws.iter_rows() for c in r]
    assert "Wipro" in texts
    link_cells = [c for r in ws.iter_rows() for c in r if c.hyperlink is not None]
    assert any(c.hyperlink.target == "https://example.com/led" for c in link_cells)


from app.services.langchain.tools.cost_calculator_tool import _CalcLine


def test_calcline_accepts_oem_and_source_url():
    ln = _CalcLine(description="X", quantity=1, rate=10,
                   oem_manufacturer="ACME", source_url="https://acme.test")
    assert ln.oem_manufacturer == "ACME"
    assert ln.source_url == "https://acme.test"


def test_prompt_documents_oem_and_source_url():
    from app.services.langchain.graphs.costing_agent import COSTING_AGENT_SYSTEM_PROMPT
    p = COSTING_AGENT_SYSTEM_PROMPT
    assert "oem_manufacturer" in p
    assert "source_url" in p


def test_margin_layout_has_clickable_web_source(tmp_path):
    """Margin/multi-sheet layout (_build_multi_sheet) must also render the Web
    Source hyperlink. Force that layout via strategic_summary + single_sheet
    (rows with item_code+schedule_name would otherwise pick the NIT-mirror
    layout, which intentionally omits these columns)."""
    import openpyxl
    from app.services.langchain.tools.xlsx_generator_tool import build_cost_xlsx
    out = tmp_path / "margin.xlsx"
    build_cost_xlsx(
        str(out), "Cost",
        rows=[{"sr_no": "1", "description": "LED", "quantity": 2,
               "tender_rate": 600, "tender_amount": 1200,
               "rate": 500, "amount": 1000, "rate_source": "web_search",
               "oem_manufacturer": "Wipro",
               "source_url": "https://example.com/led"}],
        strategic_summary={"tender_snapshot": {"tender_no": "T1"},
                           "schedule_breakdown": []},
        single_sheet=True,
    )
    wb = openpyxl.load_workbook(str(out))
    link_targets = [
        c.hyperlink.target
        for ws in wb.worksheets for row in ws.iter_rows() for c in row
        if c.hyperlink is not None
    ]
    assert "https://example.com/led" in link_targets, \
        "margin layout did not render a clickable Web Source hyperlink"


def test_web_search_results_carry_structured_source_url():
    """Each web_search result dict exposes source_url (=url) + oem_manufacturer
    so the field is structured at the tool boundary, not only in prose."""
    import json
    from app.services.langchain.tools.web_search_tool import WebSearchTool

    # DuckDuckGo path needs no API key; monkeypatch DDGS to a fixed result.
    tool = WebSearchTool()
    fake = [{"title": "LED fitting", "href": "https://shop.test/led", "body": "Wipro LED ₹500"}]

    class _FakeDDGS:
        def __enter__(self): return self
        def __exit__(self, *a): return False
        def text(self, *a, **k): return iter(fake)

    import app.services.langchain.tools.web_search_tool as wst
    orig = getattr(wst, "DDGS", None)
    try:
        # DDGS is imported inside the method; patch the module the method imports from.
        import duckduckgo_search
        duckduckgo_search.DDGS = _FakeDDGS
        raw = tool._search_duckduckgo("led", 1, None)
        payload = json.loads(raw)
        r = payload["results"][0]
        assert r["source_url"] == "https://shop.test/led"
        assert "oem_manufacturer" in r  # present (best-effort, may be None)
    finally:
        if orig is not None:
            duckduckgo_search.DDGS = orig
