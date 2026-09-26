"""Piece 2 — reconciliation surfaced in the strategic-summary xlsx + chat reply."""
import openpyxl
from app.services.langchain.tools.xlsx_generator_tool import build_cost_xlsx


def _recon_report(ok=False):
    return {
        "ok": ok,
        "schedules": [
            {"code": "A", "line_sum": 100.0, "stated_total": 100.0, "delta": 0.0, "ok": True},
            {"code": "B", "line_sum": 250.0, "stated_total": 200.0, "delta": 50.0, "ok": False},
        ],
        "grand_line_sum": 350.0,
        "advertised_value": 300.0,
        "grand_delta": 50.0,
    }


def test_multi_sheet_renders_reconciliation_block(tmp_path):
    # Strategic-summary layout is selected when strategic_summary is present.
    out = tmp_path / "cost.xlsx"
    build_cost_xlsx(
        str(out),
        "Test Cost Sheet",
        rows=[
            {"sr_no": "1", "description": "Item A", "schedule_name": "A",
             "quantity": 1, "tender_rate": 100, "tender_amount": 100,
             "rate": 90, "amount": 90},
        ],
        strategic_summary={
            "tender_snapshot": {"tender_no": "T1"},
            "schedule_breakdown": [
                {"schedule": "Schedule A", "tender_value_inr": 100,
                 "estimated_cost_inr": 90, "gross_margin_inr": 10, "gross_margin_pct": 10.0},
            ],
        },
        breakdown_meta={"tender_id": 1, "reconciliation": _recon_report(ok=False)},
    )
    wb = openpyxl.load_workbook(str(out))
    ws = wb["1. Summary"]
    texts = [c.value for r in ws.iter_rows() for c in r if isinstance(c.value, str)]
    assert any("Reconciliation" in t for t in texts), "reconciliation block missing from Summary sheet"
    # The failing schedule B must be shown with a FLAG status.
    assert any(t == "FLAG" for t in texts), "failed schedule not flagged in reconciliation block"


def test_emit_artifact_threads_reconciliation_into_meta(monkeypatch):
    """_emit_costing_xlsx_artifact must put the persisted reconciliation report
    into breakdown_meta so the multi-sheet layout can render it. We capture the
    breakdown_meta handed to build_cost_xlsx."""
    import app.services.langchain.graphs.chat_agent_wrappers as caw

    captured = {}

    def _fake_build(path, title, rows, *, cost_assumptions, strategic_summary,
                    breakdown_meta, single_sheet=False):
        captured["meta"] = breakdown_meta
        # Write a tiny valid xlsx so the surrounding code can read bytes.
        import openpyxl
        openpyxl.Workbook().save(path)
        return {"total_amount": 0, "total_with_gst": 0}

    monkeypatch.setattr(caw, "build_cost_xlsx", _fake_build)

    # Stub the breakdown lookup to return a report.
    report = {"ok": False, "schedules": [], "grand_line_sum": 0,
              "advertised_value": None, "grand_delta": 0}
    monkeypatch.setattr(caw, "_latest_reconciliation_for_tender",
                        lambda db, tid: report)

    # Stub the storage upload and artifact-creation steps that run *after*
    # build_cost_xlsx — this test only cares about the breakdown_meta wiring,
    # not real persistence/storage side effects. _emit_costing_xlsx_artifact
    # imports these lazily inside its own body, so patch at their source
    # modules (a patch on `caw` would not affect the fresh local import).
    class _FakeStorage:
        def upload_file_sync(self, key, data, content_type=None):
            pass

    import app.services.storage_service as storage_service
    monkeypatch.setattr(storage_service, "get_storage_service", lambda: _FakeStorage())

    class _FakeArtifact:
        id = 1
        version = 1

    import app.services.artifact_service as artifact_service
    monkeypatch.setattr(
        artifact_service, "create_artifact", lambda **kwargs: _FakeArtifact()
    )

    # Minimal stubs for the rest of the emit path. _emit_costing_xlsx_artifact
    # returns None early when there are no line items, so include one — the
    # test cares about breakdown_meta wiring, not the row content.
    class _DB:
        def add(self, *a, **kw): pass
        def commit(self, *a, **kw): pass
        def query(self, *a, **kw):
            raise AssertionError("unexpected DB query in wiring test")

    costing = {
        "line_items": [{"description": "Item A", "quantity": 1, "rate": 10, "amount": 10}],
        "strategic_summary": {},
        "cost_assumptions": [],
    }
    caw._emit_costing_xlsx_artifact(_DB(), proposal_session_id=None,
                                    tender_id=1, costing=costing)
    assert captured["meta"].get("reconciliation") == report


from app.services.langchain.graphs.chat_agent_wrappers import _reconciliation_callout


def test_callout_none_when_reconciled():
    report = {"ok": True, "schedules": [
        {"code": "A", "line_sum": 100, "stated_total": 100, "delta": 0, "ok": True}]}
    assert _reconciliation_callout(report) is None


def test_callout_none_when_no_report():
    assert _reconciliation_callout(None) is None


def test_callout_warns_and_counts_failed_schedules():
    report = {"ok": False, "schedules": [
        {"code": "A", "line_sum": 100, "stated_total": 100, "delta": 0, "ok": True},
        {"code": "B", "line_sum": 250, "stated_total": 200, "delta": 50, "ok": False},
        {"code": "C", "line_sum": 300, "stated_total": 250, "delta": 50, "ok": False}]}
    out = _reconciliation_callout(report)
    assert out is not None
    assert "⚠️" in out
    assert "2" in out          # two schedules failed
    assert "B" in out and "C" in out  # names the failing schedules


def _has_reconciliation(path):
    wb = openpyxl.load_workbook(str(path))
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and "Reconciliation" in c.value:
                    return True
    return False


def test_single_sheet_layout_renders_reconciliation_block(tmp_path):
    # Legacy single-sheet layout (no strategic_summary, no margin data).
    out = tmp_path / "single.xlsx"
    build_cost_xlsx(
        str(out), "Cost",
        rows=[{"sr_no": "1", "description": "Item", "qty": 1, "rate": 10, "amount": 10}],
        breakdown_meta={"tender_id": 1, "reconciliation": _recon_report(ok=False)},
    )
    assert _has_reconciliation(out), "single-sheet layout omitted reconciliation block"


def test_client_annexure_layout_renders_reconciliation_block(tmp_path):
    # Client-annexure layout is selected when rows carry an `annexure` field.
    out = tmp_path / "ann.xlsx"
    build_cost_xlsx(
        str(out), "Cost",
        rows=[{"sr_no": "1", "item_code": "P1", "description": "Part", "annexure": "A",
               "quantity": 2, "rate": 50, "amount": 100}],
        breakdown_meta={"tender_id": 1, "reconciliation": _recon_report(ok=False)},
    )
    assert _has_reconciliation(out), "client-annexure layout omitted reconciliation block"
